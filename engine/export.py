"""Excel export of the P&L tree and the scenario levers."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import simulate as sim

LEVER_LABELS = {
    "price_change_pct": ("Price (MRP) change", "%"),
    "promo_add_pp": ("Extra promo discount", "pp of invoice"),
    "promo_pass_through_pct": ("Promo passed to shelf price", "%"),
    "trade_margin_add_pp": ("Trade margin change", "pp of MRP"),
    "rebate_add_pp": ("Other deductions and display fees", "pp of invoice"),
    "elasticity_scale": ("Price sensitivity multiplier", "x"),
    "manual_uplift_pct": ("Extra volume from visibility", "%"),
    "cogs_change_pct": ("COGS change", "%"),
    "freight_change_pct": ("S&D / freight cost", "%"),
}

HEADER_FILL = PatternFill("solid", fgColor="123524")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FONT = Font(bold=True)

# Negative numbers show in red with a minus sign.
AMOUNT = "#,##0.0;[Red]-#,##0.0"
PERCENT = "0.0;[Red]-0.0"

# (header, metric key, scale, number format)
COLUMNS = [
    ("Volume ('000 units)", "units", 1e3, AMOUNT),
    ("Gross sales", "gross", 1e5, AMOUNT),
    ("Promo", "promo", 1e5, AMOUNT),
    ("Other deductions", "rebates", 1e5, AMOUNT),
    ("Net revenue", "net", 1e5, AMOUNT),
    ("COGS", "cogs", 1e5, AMOUNT),
    ("Gross margin", "gm", 1e5, AMOUNT),
    ("Over/under", "ou", 1e5, AMOUNT),
    ("S&D / freight", "logistics", 1e5, AMOUNT),
    ("Contribution", "cm", 1e5, AMOUNT),
    ("P&A", "pa", 1e5, AMOUNT),
    ("Interest", "interest", 1e5, AMOUNT),
    ("Marketing", "mktg", 1e5, AMOUNT),
    ("Net margin", "nm", 1e5, AMOUNT),
    ("GM %", "gm_pct", 0.01, PERCENT),
    ("CM %", "cm_pct", 0.01, PERCENT),
    ("NM %", "nm_pct", 0.01, PERCENT),
]


def _flatten(nodes, depth=0):
    for n in nodes:
        yield depth, n
        yield from _flatten(n.get("children", []), depth + 1)


def _filter_text(flt_raw):
    from . import data
    parts = []
    for dim, title in data.DIM_DEFS:
        ids = (flt_raw or {}).get(dim) or []
        names = [data.LABEL[dim][i] for i in ids if i in data.LABEL[dim]]
        if names:
            parts.append(f"{title}: {', '.join(names)}")
    return "; ".join(parts) or "none (all rows)"


def _basis_text():
    from . import data
    months = data.months_text(data.META.get("base_months", []))
    if data.MODE == "actual":
        return f"Monthly run-rate, {months} actual billing." if months else "Monthly run-rate, actual billing."
    return "Monthly run-rate. Sample data, not actuals."


PIVOT_SHEET = "Data for pivot"
PIVOT_STEPS = [
    "To build a PivotTable like the MIS one:",
    "1. Click any cell in the table on the left.",
    "2. Insert > PivotTable > OK.",
    "3. Drag Chain, State, City, Product category, SKU (or any you like) into Rows, in the order you want.",
    "4. Drag the amounts you want into Values (for example Scenario Net revenue, Scenario Contribution). Choose Sum.",
    "5. For a percentage, use PivotTable Analyze > Fields, Items & Sets > Calculated Field.",
    "   Example, name CM %: = 'Scenario Contribution' / 'Scenario Net revenue'. Format it as a percentage.",
    "Amounts are Rs lakh a month, volume is thousand units. Baseline columns have no levers applied.",
]


def _pivot_sheet(wb, levers_raw, scope_raw, filter_raw):
    """A flat sheet, one line per dataset row, that a PivotTable can be built on in a few clicks."""
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from . import data

    items = sim.pnl_items(levers_raw, scope_raw, filter_raw)
    ws = wb.create_sheet(PIVOT_SHEET)
    metric_cols = [c for c in COLUMNS if not c[1].endswith("_pct")]
    header = [title for _, title in data.DIM_DEFS]
    header += [f"{side} {name}" for side in ("Baseline", "Scenario") for name, _, _, _ in metric_cols]
    ws.append(header)
    for it in items:
        row = [data.LABEL[d].get(it["row"][d], it["row"][d]) for d, _ in data.DIM_DEFS]
        for side in ("base", "sim"):
            row += [it[side][k] / scale for _, k, scale, _ in metric_cols]
        ws.append(row)
    n_dims = len(data.DIM_DEFS)
    for r in range(2, ws.max_row + 1):
        for c in range(n_dims + 1, len(header) + 1):
            ws.cell(row=r, column=c).number_format = AMOUNT
    if ws.max_row > 1:
        last = f"{get_column_letter(len(header))}{ws.max_row}"
        table = Table(displayName="PnlRows", ref=f"A1:{last}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    for i in range(1, n_dims + 1):
        ws.column_dimensions[get_column_letter(i)].width = 22
    for i in range(n_dims + 1, len(header) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 17
    ws.freeze_panes = "A2"
    note_col = len(header) + 2
    for k, line in enumerate(PIVOT_STEPS, start=1):
        cell = ws.cell(row=k, column=note_col, value=line)
        cell.font = TOTAL_FONT if k == 1 else Font()
    ws.column_dimensions[get_column_letter(note_col)].width = 70


def build_workbook(levers_raw, scope_raw, dims, filter_raw=None):
    res = sim.pnl(levers_raw, scope_raw, dims, filter_raw)
    lv = sim.parse_levers(levers_raw)
    wb = Workbook()

    for title, side in (("Scenario P&L", "sim"), ("Baseline P&L", "base")):
        ws = wb.active if title == "Scenario P&L" else wb.create_sheet(title)
        ws.title = title
        ws.append([f"Amounts in Rs lakh unless stated. {_basis_text()} "
                   f"Filters: {_filter_text(filter_raw)}"])
        ws.append(["Name"] + [c[0] for c in COLUMNS])
        for cell in ws[2]:
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

        def emit(label, metrics, depth, bold=False):
            ws.append(["    " * depth + label] +
                      [metrics[k] / scale for _, k, scale, _ in COLUMNS])
            row = ws.max_row
            for i, (_, _, _, fmt) in enumerate(COLUMNS, start=2):
                ws.cell(row=row, column=i).number_format = fmt
            if bold:
                for cell in ws[row]:
                    cell.font = TOTAL_FONT
            # Excel's own grouping: the + / - buttons and the 1, 2, 3 level buttons at the top left.
            # The Total row and the first level stay at outline level 0.
            ws.row_dimensions[row].outlineLevel = depth

        emit("Total", res["total"][side], 0, bold=True)
        for depth, node in _flatten(res["tree"]):
            emit(node["label"], node[side], depth, bold=depth == 0)
        ws.sheet_properties.outlinePr.summaryBelow = False   # the + / - button sits on the parent row, above its children
        ws.column_dimensions["A"].width = 44
        for i in range(2, len(COLUMNS) + 2):
            ws.column_dimensions[get_column_letter(i)].width = 15
        ws.row_dimensions[2].height = 32
        ws.freeze_panes = "B3"

    _pivot_sheet(wb, levers_raw, scope_raw, filter_raw)

    ws = wb.create_sheet("Scenario levers")
    ws.append(["Lever", "Value", "Unit"])
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for key, (label, unit) in LEVER_LABELS.items():
        ws.append([label, getattr(lv, key), unit])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["C"].width = 16

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
