import math

import pytest

from engine import data, simulate as sim


def first(dim):
    return data.OPTIONS[dim][0]["id"]


def last(dim):
    return data.OPTIONS[dim][-1]["id"]


def close(a, b, tol=1e-6):
    return math.isclose(a, b, rel_tol=tol, abs_tol=1e-6)


def test_tests_run_on_sample_data():
    # guards against a test run quietly using real company data
    assert data.MODE == "sample"


def test_neutral_levers_match_baseline():
    out = sim.simulate({}, {})
    for k in sim.METRICS:
        assert close(out["base"][k], out["sim"][k])
    assert all(abs(s["value"]) < 1e-6 for s in out["bridge"] if s["kind"] == "delta")


def test_pnl_lines_add_up():
    m = sim.simulate({}, {})["base"]
    assert close(m["gross"], m["mrp_value"] - m["trade_margin"])
    assert close(m["net"], m["gross"] - m["promo"] - m["rebates"])
    assert close(m["gm"], m["net"] - m["cogs"])
    assert close(m["cm"], m["gm"] - m["ou"] - m["logistics"])          # the MIS definition of contribution
    assert close(m["nm"], m["cm"] - m["pa"] - m["interest"] - m["mktg"])


def test_over_under_sits_above_contribution_and_net_margin_is_unchanged():
    m = sim.simulate({}, {})["base"]
    assert m["cm"] < m["gm"] - m["logistics"]                           # over/under now reduces contribution
    old_nm = (m["gm"] - m["logistics"]) - m["pa"] - m["interest"] - m["ou"] - m["mktg"]
    assert close(m["nm"], old_nm)                                       # the bottom line does not move


def test_every_row_of_the_pnl_follows_the_same_definition():
    out = sim.pnl({"promo_add_pp": 3}, {}, ["chain", "state"], {})
    seen = 0

    def walk(nodes):
        nonlocal seen
        for n in nodes:
            for side in ("base", "sim"):
                m = n[side]
                assert close(m["cm"], m["gm"] - m["ou"] - m["logistics"])
                assert close(m["nm"], m["cm"] - m["pa"] - m["interest"] - m["mktg"])
                seen += 1
            walk(n.get("children") or [])

    walk(out["tree"])
    assert seen > 4


def test_price_rise_raises_realisation_and_leaves_volume_to_you():
    out = sim.simulate({"price_change_pct": 5}, {})
    assert close(out["sim"]["units"], out["base"]["units"])      # volume only moves when you enter a change
    assert out["sim"]["net_per_unit"] > out["base"]["net_per_unit"]
    lower = sim.simulate({"price_change_pct": 5, "volume_change_pct": -4}, {})
    assert lower["sim"]["units"] < lower["base"]["units"]


def test_promo_costs_margin_rate_and_adds_volume_when_you_expect_it():
    out = sim.simulate({"promo_add_pp": 5, "volume_change_pct": 8}, {})
    assert out["sim"]["units"] > out["base"]["units"]
    assert out["sim"]["promo"] > out["base"]["promo"]
    assert out["sim"]["gm_pct"] < out["base"]["gm_pct"]


def test_attributed_costs_follow_their_basis():
    # On one line, volume-based costs move with packs sold and revenue-based costs move with net revenue.
    # (Across many lines the ratios also shift with mix, so this is checked line by line.)
    row = max(data.ROWS, key=lambda r: r["units"])
    before = sim.calc_row(row, sim.DEFAULT_LEVERS, False)
    after = sim.calc_row(row, sim.parse_levers({"price_change_pct": 5, "volume_change_pct": 5}), True)
    assert after["units"] != before["units"]
    for line in ("cogs", "sd", "ou"):
        assert data.BASIS[line] == "volume"
    for line in ("pa", "interest", "mktg"):
        assert data.BASIS[line] == "revenue"
        assert close(after[line] / after["net"], before[line] / before["net"])
    assert close(after["ou"] / after["units"], before["ou"] / before["units"])


def test_cogs_change_hits_cogs_everywhere_even_with_scope():
    out = sim.simulate({"cogs_change_pct": 10}, {"chain": [first("chain")]})
    other = [c for c in out["breakdown"]["chain"] if c["id"] == last("chain")][0]
    assert out["sim"]["cogs"] > out["base"]["cogs"]
    assert other["sim"]["cogs"] > other["base"]["cogs"]
    assert out["sim"]["cogs"] == pytest.approx(out["base"]["cogs"] * 1.10)      # exactly +10%, no hidden share


def test_scope_limits_commercial_levers():
    c1, c2 = first("chain"), last("chain")
    out = sim.simulate({"price_change_pct": 8}, {"chain": [c1]})
    by = {c["id"]: c for c in out["breakdown"]["chain"]}
    assert not close(by[c1]["sim"]["net"], by[c1]["base"]["net"])
    assert close(by[c2]["sim"]["net"], by[c2]["base"]["net"])


@pytest.mark.parametrize("dim", data.DIM_IDS)
def test_scope_works_on_every_dimension(dim):
    out = sim.simulate({"price_change_pct": 4, "volume_change_pct": 5}, {dim: [first(dim)]})
    assert 0 < out["rows_in_scope"] <= out["rows_total"]
    if len(data.OPTIONS[dim]) > 1:
        assert out["rows_in_scope"] < out["rows_total"]
        assert not close(out["sim"]["units"], out["base"]["units"])


def test_unknown_scope_id_changes_nothing():
    out = sim.simulate({"price_change_pct": 8}, {"chain": ["no-such-chain"]})
    assert out["rows_in_scope"] == 0
    assert close(out["sim"]["net"], out["base"]["net"])


def test_bridge_sums_to_scenario():
    out = sim.simulate({"price_change_pct": 3, "promo_add_pp": 4,
                        "trade_margin_add_pp": 1, "cogs_change_pct": 5}, {})
    steps = out["bridge"]
    assert close(steps[0]["value"] + sum(s["value"] for s in steps[1:-1]), steps[-1]["value"])
    assert close(steps[-1]["value"], out["sim"]["cm"])


def test_tree_children_sum_to_parent():
    res = sim.pnl({"promo_add_pp": 3}, {}, ["chain", "state", "city", "sku"])
    assert res["dims"] == ["chain", "state", "city", "sku"]
    for node in res["tree"]:
        for key in ("net", "cm", "nm", "units"):
            assert close(node["sim"][key], sum(c["sim"][key] for c in node["children"]))
    for key in ("net", "cm", "nm"):
        assert close(res["total"]["sim"][key], sum(n["sim"][key] for n in res["tree"]))


def test_every_dimension_can_group_the_pnl():
    for dim in data.DIM_IDS:
        res = sim.pnl({}, {}, [dim])
        assert res["tree"], dim
        assert close(res["total"]["base"]["net"], sum(n["base"]["net"] for n in res["tree"]))


def test_inputs_are_clamped():
    lv = sim.parse_levers({"price_change_pct": 999, "cogs_change_pct": "abc"})
    assert lv.price_change_pct == sim.LIMITS["price_change_pct"][1]
    assert lv.cogs_change_pct == 0.0


def test_pnl_filter_limits_rows_and_totals():
    full = sim.pnl({}, {}, ["chain"])
    part = sim.pnl({}, {}, ["chain"], {"city": [first("city")]})
    assert 0 < part["rows_shown"] < full["rows_shown"] == part["rows_total"]
    assert part["total"]["sim"]["net"] < full["total"]["sim"]["net"]
    assert close(part["total"]["sim"]["net"], sum(n["sim"]["net"] for n in part["tree"]))


def test_pnl_filter_does_not_change_scenario_per_row():
    # filtering is a view: a state's numbers are the same with or without a state filter
    lv = {"price_change_pct": 3}
    state = first("state")
    a = sim.pnl(lv, {}, ["state"], {"state": [state]})["tree"][0]["sim"]["net"]
    b = [n for n in sim.pnl(lv, {}, ["state"])["tree"] if n["id"].endswith(state)][0]["sim"]["net"]
    assert close(a, b)


def test_filters_on_several_dimensions_combine():
    one = sim.pnl({}, {}, ["chain"], {"grp": [first("grp")]})["rows_shown"]
    both = sim.pnl({}, {}, ["chain"], {"grp": [first("grp")], "state": [first("state")]})["rows_shown"]
    assert 0 < both <= one


def test_old_style_plural_keys_are_ignored_not_misread():
    # an old client sending {"chains": [...]} must not silently filter on something else
    out = sim.simulate({"price_change_pct": 5}, {"chains": [first("chain")]})
    assert out["rows_in_scope"] == out["rows_total"]


# ---------- breakeven and the volume ladder ----------
def base_cm():
    return sum(m["cm"] for m in sim.BASE_METRICS)


def cm_with_uplift(levers, scope, uplift):
    """Contribution if every pack in scope rises by `uplift`, built row by row at the scenario's terms.
    Independent of breakeven(): it only uses calc_row."""
    lv, sc = sim.parse_levers(levers), sim.parse_scope(scope)
    total = 0.0
    for r in data.ROWS:
        active = sim.in_scope(r, sc)
        row = dict(r, units=r["units"] * (1 + uplift)) if active else r
        total += sim.calc_row(row, lv, active, fixed_volume=True)["cm"]
    return total


def test_breakeven_is_neutral_when_nothing_changes():
    b = sim.simulate({}, {})["breakeven"]
    assert b["status"] == "neutral" and b["required"] is None


def test_extra_volume_alone_has_nothing_to_pay_back():
    # it adds volume but changes no price, promo or cost, so breakeven has nothing to say
    out = sim.simulate({"volume_change_pct": 6}, {})
    assert out["breakeven"]["status"] == "neutral"
    assert out["sim"]["units"] > out["base"]["units"]


@pytest.mark.parametrize("levers", [
    {"promo_add_pp": 5},
    {"price_change_pct": 3},
    {"cogs_change_pct": 5},
    {"trade_margin_add_pp": 2, "rebate_add_pp": 1},
    {"promo_add_pp": 3, "price_change_pct": 2, "freight_change_pct": 10},
])
@pytest.mark.parametrize("scope", [{}, "one chain"])
def test_breakeven_volume_brings_contribution_back_to_baseline(levers, scope):
    scope = {"chain": [first("chain")]} if scope == "one chain" else scope
    b = sim.simulate(levers, scope)["breakeven"]
    assert b["status"] in ("needs", "ahead")
    assert close(cm_with_uplift(levers, scope, b["required"]), base_cm(), tol=1e-9)


def test_a_promo_needs_extra_volume_and_you_may_not_expect_enough():
    b = sim.simulate({"promo_add_pp": 5}, {})["breakeven"]
    assert b["status"] == "needs" and b["required"] > 0 and b["extra_units"] > 0
    assert b["covered"] is False and b["gap"] < 0
    assert close(b["gap"], b["assumed"] - b["required"])


def test_a_price_rise_is_ahead_and_can_lose_volume():
    b = sim.simulate({"price_change_pct": 3, "volume_change_pct": -2}, {})["breakeven"]
    assert b["status"] == "ahead" and b["required"] < 0
    assert close(b["assumed"], -0.02) and b["covered"] is True


def test_breakeven_is_unreachable_when_each_extra_pack_loses_money():
    b = sim.simulate({"promo_add_pp": 20, "trade_margin_add_pp": 8, "rebate_add_pp": 6}, {})["breakeven"]
    assert b["status"] == "unreachable" and b["required"] is None


def test_breakeven_counts_only_the_lines_in_scope():
    chain = first("chain")
    b = sim.simulate({"promo_add_pp": 5}, {"chain": [chain]})["breakeven"]
    in_scope_units = sum(r["units"] for r in data.ROWS if r["chain"] == chain)
    all_units = sum(r["units"] for r in data.ROWS)
    assert close(b["units_in_scope"], in_scope_units) and b["units_in_scope"] < all_units


def test_breakeven_with_nothing_in_scope_is_neutral():
    b = sim.simulate({"promo_add_pp": 5}, {"chain": ["no-such-chain"]})["breakeven"]
    assert b["status"] == "neutral" and b["units_in_scope"] == 0


def test_the_volume_you_enter_scales_units_in_scope_only():
    chain = first("chain")
    out = sim.simulate({"volume_change_pct": 10}, {"chain": [chain]})
    inside = sum(r["units"] for r in data.ROWS if r["chain"] == chain)
    assert close(out["sim"]["units"] - out["base"]["units"], 0.10 * inside)
    assert close(out["breakeven"]["assumed"], 0.10)


def test_volume_alone_changes_no_rate():
    out = sim.simulate({"volume_change_pct": 10}, {})
    assert close(out["sim"]["units"], 1.10 * out["base"]["units"])
    assert close(out["sim"]["net"], 1.10 * out["base"]["net"])
    assert close(out["sim"]["gross"] / out["sim"]["units"], out["base"]["gross"] / out["base"]["units"])


def test_default_volume_is_zero_so_levers_alone_do_not_move_volume():
    out = sim.simulate({"promo_add_pp": 5, "price_change_pct": -3}, {})
    assert close(out["sim"]["units"], out["base"]["units"])


def test_the_ladder_marks_the_entered_volume_and_matches_the_scenario():
    out = sim.simulate({"promo_add_pp": 3, "volume_change_pct": 7}, {})
    cases = out["ladder"]["cases"]
    entered = [c for c in cases if c["entered"]]
    assert len(entered) == 1 and entered[0]["volume_pct"] == 7
    assert close(entered[0]["cm"], out["sim"]["cm"]) and close(entered[0]["units"], out["sim"]["units"])
    assert [c["volume_pct"] for c in cases] == sorted({-10, 0, 10, 20, 30, 7})


def test_the_ladder_has_no_duplicate_when_the_entry_is_a_step():
    cases = sim.simulate({"promo_add_pp": 3, "volume_change_pct": 10}, {})["ladder"]["cases"]
    assert [c["volume_pct"] for c in cases] == [-10, 0, 10, 20, 30]
    assert sum(c["entered"] for c in cases) == 1


def test_the_ladder_crosses_zero_at_the_breakeven_volume():
    levers = {"promo_add_pp": 2}
    out = sim.simulate(levers, {})
    required = out["breakeven"]["required"] * 100
    at = sim.simulate({**levers, "volume_change_pct": required}, {})
    assert close(at["sim"]["cm"], at["base"]["cm"], tol=1e-9)


def test_ladder_contribution_rises_with_volume_when_a_pack_earns_money():
    cms = [c["cm"] for c in sim.simulate({"promo_add_pp": 2}, {})["ladder"]["cases"]]
    assert cms == sorted(cms)


def test_the_model_suggests_nothing_when_no_lever_is_pulled():
    assert abs(sim.simulate({}, {})["suggested_volume"]) < 1e-12


def test_the_suggestion_follows_the_direction_of_the_levers():
    assert sim.simulate({"promo_add_pp": 4}, {})["suggested_volume"] > 0
    assert sim.simulate({"price_change_pct": 4}, {})["suggested_volume"] < 0


def test_the_suggestion_never_changes_the_numbers():
    a = sim.simulate({"promo_add_pp": 4}, {})
    assert a["suggested_volume"] != 0 and close(a["sim"]["units"], a["base"]["units"])


def test_cost_levers_do_not_depend_on_the_scope_of_volume():
    chain = first("chain")
    whole = sim.simulate({"cogs_change_pct": 5}, {})
    scoped = sim.simulate({"cogs_change_pct": 5}, {"chain": [chain]})
    assert close(whole["sim"]["cogs"], scoped["sim"]["cogs"])


def test_a_large_fall_never_makes_volume_negative():
    lv = sim.parse_levers({"volume_change_pct": -1000})
    assert lv.volume_change_pct == sim.LIMITS["volume_change_pct"][0]
    assert all(sim.calc_row(r, lv, True)["units"] >= 0 for r in data.ROWS)


def test_the_bridge_volume_step_is_the_volume_you_entered():
    out = sim.simulate({"promo_add_pp": 3, "volume_change_pct": 8}, {})
    steps = {s["key"]: s for s in out["bridge"]}
    assert steps["volume"]["label"] == "Volume you expect"
    assert steps["volume"]["value"] > 0
    assert close(out["bridge"][0]["value"] + sum(s["value"] for s in out["bridge"][1:-1]), out["sim"]["cm"])


def test_excel_export_has_collapsible_rows_and_red_negatives():
    import openpyxl
    from engine import export
    buf = export.build_workbook({"promo_add_pp": 3}, {}, ["chain", "state", "city"], {})
    wb = openpyxl.load_workbook(buf)
    for name in ("Scenario P&L", "Baseline P&L"):
        ws = wb[name]
        assert ws.sheet_properties.outlinePr.summaryBelow is False
        levels = {}
        for r in range(3, ws.max_row + 1):
            label = ws.cell(row=r, column=1).value
            depth = (len(label) - len(label.lstrip())) // 4
            levels.setdefault(depth, set()).add(ws.row_dimensions[r].outlineLevel)
        # every row's outline level equals its indent, so Excel can collapse each level
        assert all(v == {d} for d, v in levels.items())
        assert max(levels) == 2
        # negatives are red, amounts and percentages alike
        assert "[Red]" in ws.cell(row=3, column=2).number_format
        assert "[Red]" in ws.cell(row=3, column=ws.max_column).number_format


def test_excel_export_has_a_flat_sheet_that_adds_up_to_the_total_row():
    import openpyxl
    from engine import export
    levers = {"promo_add_pp": 3}
    flt = {"chain": [first("chain")]}
    wb = openpyxl.load_workbook(export.build_workbook(levers, {}, ["chain", "state"], flt))
    ws = wb[export.PIVOT_SHEET]
    head = [c.value for c in ws[1]][:len(data.DIM_DEFS) + 2 * 14]
    assert head[:len(data.DIM_DEFS)] == [t for _, t in data.DIM_DEFS]
    assert len(ws.tables) == 1
    assert len(set(head)) == len(head)                      # a table needs unique header names
    i_net = head.index("Scenario Net revenue") + 1
    i_base = head.index("Baseline Net revenue") + 1
    flat_sim = sum(ws.cell(row=r, column=i_net).value for r in range(2, ws.max_row + 1)
                   if ws.cell(row=r, column=1).value is not None)
    flat_base = sum(ws.cell(row=r, column=i_base).value for r in range(2, ws.max_row + 1)
                    if ws.cell(row=r, column=1).value is not None)
    total = sim.pnl(levers, {}, ["chain"], flt)["total"]
    assert flat_sim == pytest.approx(total["sim"]["net"] / 1e5)
    assert flat_base == pytest.approx(total["base"]["net"] / 1e5)
    assert ws.max_row - 1 == sim.pnl(levers, {}, ["chain"], flt)["rows_shown"]
